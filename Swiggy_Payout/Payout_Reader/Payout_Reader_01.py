import pandas as pd
from pathlib import Path
import re

all_rows = []

for file in Path(".").glob("*.xlsx"):

    if file.name.startswith("~$"):
        continue

    try:
        print(f"\nProcessing: {file.name}")

        # =====================================
        # READ SUMMARY SHEET
        # =====================================
        restaurant_id = ""
        payout_period = ""
        payout_settlement_date = ""

        try:
            summary = pd.read_excel(
                file,
                sheet_name="Summary",
                header=None
            ).fillna("")

            for r in range(summary.shape[0]):
                for c in range(summary.shape[1]):

                    value = str(summary.iloc[r, c]).strip()

                    # Restaurant ID
                    if "Rest. ID" in value:

                        match = re.search(r"(\d+)", value)

                        if match:
                            restaurant_id = match.group(1)

                    # Payout Period
                    if value.lower() == "payout period":

                        if c + 1 < summary.shape[1]:
                            payout_period = str(
                                summary.iloc[r, c + 1]
                            ).strip()

                    # Settlement Date
                    if value.lower() == "payout settlement date":

                        if c + 1 < summary.shape[1]:
                            payout_settlement_date = str(
                                summary.iloc[r, c + 1]
                            ).strip()

        except Exception as e:
            print(f"Summary Sheet Error: {e}")

        # =====================================
        # READ ORDER LEVEL SHEET
        # =====================================
        raw = pd.read_excel(
            file,
            sheet_name="Order Level",
            header=None
        )

        header_row = None

        for i in range(min(30, len(raw))):

            row_text = " ".join(
                str(x).strip()
                for x in raw.iloc[i].fillna("").tolist()
            )

            if "Order Status" in row_text:
                header_row = i
                break

        if header_row is None:
            print("Order Status header not found")
            continue

        df = pd.read_excel(
            file,
            sheet_name="Order Level",
            header=header_row
        )

        df.columns = (
            df.columns.astype(str)
            .str.strip()
            .str.replace("\n", " ", regex=False)
        )

        status_col = None

        for col in df.columns:
            if str(col).strip().lower() == "order status":
                status_col = col
                break

        if status_col is None:
            print("Order Status column not found")
            continue

        df[status_col] = (
            df[status_col]
            .astype(str)
            .str.strip()
            .str.lower()
        )

        delivered = df[
            df[status_col] == "delivered"
        ]

        cancelled = df[
            df[status_col] == "cancelled"
        ]

        # =====================================
        # INVOICE NUMBER
        # =====================================
        invoice_match = re.search(
            r"invoice_Annexure_(\d+)",
            file.name,
            re.IGNORECASE
        )

        invoice_no = (
            invoice_match.group(1)
            if invoice_match
            else file.stem
        )

        # =====================================
        # BASE ROW
        # =====================================
        row = {
            "Restaurant ID": restaurant_id,
            "Payout Period": payout_period,
            "Payout Settlement Date": payout_settlement_date,
            "Invoice No": invoice_no,
            "File Name": file.name,
            "Delivered Orders": len(delivered),
            "Cancelled Orders": len(cancelled)
        }

        # =====================================
        # NUMERIC COLUMNS
        # =====================================
        numeric_cols = []

        for col in df.columns:

            if col == status_col:
                continue

            converted = pd.to_numeric(
                df[col],
                errors="coerce"
            )

            if converted.notna().sum() > 0:
                numeric_cols.append(col)

        # =====================================
        # D_ / C_ TOTALS
        # =====================================
        for col in numeric_cols:

            d_sum = pd.to_numeric(
                delivered[col],
                errors="coerce"
            ).sum()

            c_sum = pd.to_numeric(
                cancelled[col],
                errors="coerce"
            ).sum()

            row[f"D_{col}"] = round(d_sum, 2)
            row[f"C_{col}"] = round(c_sum, 2)

        all_rows.append(row)

        print(
            f"✓ Delivered={len(delivered)} "
            f"Cancelled={len(cancelled)}"
        )

    except Exception as e:
        print(f"Error in {file.name}: {e}")

# =====================================
# EXPORT
# =====================================
if all_rows:

    summary_df = pd.DataFrame(all_rows)

    output_file = "Consolidated_Invoice_Summary.xlsx"

    summary_df.to_excel(
        output_file,
        index=False
    )

    # =====================================
    # FORMATTING
    # =====================================
    from openpyxl import load_workbook
    from openpyxl.styles import (
        PatternFill,
        Font,
        Alignment,
        Border,
        Side
    )

    wb = load_workbook(output_file)
    ws = wb.active

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        color="FFFFFF",
        bold=True
    )

    green_fill = PatternFill(
        "solid",
        fgColor="E2F0D9"
    )

    red_fill = PatternFill(
        "solid",
        fgColor="FCE4D6"
    )

    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    # Header formatting
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    # D_ / C_ color formatting
    for column in ws.iter_cols():

        header = str(column[0].value)

        if header.startswith("D_"):

            for cell in column[1:]:
                cell.fill = green_fill

        elif header.startswith("C_"):

            for cell in column[1:]:
                cell.fill = red_fill

    # Auto width
    for column in ws.columns:

        max_len = 0

        for cell in column:

            try:
                if cell.value:
                    max_len = max(
                        max_len,
                        len(str(cell.value))
                    )
            except:
                pass

        ws.column_dimensions[
            column[0].column_letter
        ].width = min(max_len + 3, 45)

    # Freeze header
    ws.freeze_panes = "A2"

    # Filter
    ws.auto_filter.ref = ws.dimensions

    wb.save(output_file)

    print("\n====================================")
    print("DONE")
    print(f"Output File: {output_file}")
    print(f"Invoices Processed: {len(summary_df)}")
    print("====================================")

else:
    print("No invoice data found.")